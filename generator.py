\
from pathlib import Path
import re
from openpyxl import load_workbook
from .models import ContractData
from .rules import canvas_rows

TEMPLATE_PATH = Path(__file__).resolve().parent.parent / "templates" / "Webbeds_Hotel_Canvas.xlsx"


def safe_filename(name: str):
    name = re.sub(r"[^A-Za-z0-9À-ž_-]+", "_", name).strip("_")
    return f"{name}_Canva.xlsx"


def fill_canvas(contract: ContractData, output_path: str):
    wb = load_workbook(TEMPLATE_PATH)
    ws = wb["Hotel Canvas"]

    ws["C2"] = contract.hotel_name
    ws["C5"] = contract.city
    ws["C7"] = contract.country
    ws["C8"] = contract.address
    ws["C10"] = contract.latitude
    ws["C11"] = contract.longitude
    ws["C19"] = contract.website

    if contract.adult_only:
        ws["C22"] = None
        ws["C23"] = None
        ws["C24"] = None
    else:
        ws["C22"] = contract.baby_age
        ws["C23"] = contract.child1_age
        ws["C24"] = contract.child2_age

    # Category A26:A35
    for r in range(26, 36):
        ws[f"A{r}"] = None
    category_row = {
        1: 27,
        2: 28,
        3: 29,
        4: 30,
        5: 31,
        6: 32,
    }.get(contract.category)
    if category_row:
        ws[f"A{category_row}"] = "X"

    # Accommodation Type => General
    for r in range(37, 42):
        ws[f"A{r}"] = None
    ws["A40"] = "X"

    # Clear room section without unmerging A:C.
    for r in range(173, 211):
        ws[f"A{r}"] = None
        for c in range(4, 10):
            ws.cell(r, c).value = None

    rows = []
    for room in contract.rooms:
        rows.extend(canvas_rows(room, contract.adult_only))

    if len(rows) > 38:
        raise ValueError(f"Canvas room section supports 38 rows; generated {len(rows)} rows.")

    for idx, row in enumerate(rows, start=173):
        ws[f"A{idx}"] = row["name"]
        ws[f"D{idx}"] = row["min_adults"]
        ws[f"E{idx}"] = row["max_adults"]
        ws[f"F{idx}"] = row["min_children"]
        ws[f"G{idx}"] = row["max_children"]
        ws[f"H{idx}"] = None  # Always blank
        ws[f"I{idx}"] = row["max_occupancy"]

    ws["A213"] = contract.currency
    ws["A218"] = " / ".join(contract.meal_plans)

    wb.save(output_path)
    return rows
